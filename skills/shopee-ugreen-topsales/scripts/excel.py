"""Create the UGREEN Top Sales skill's single Excel deliverable.

The module deliberately accepts duck-typed scraper records.  This keeps the
workbook writer independent from the browser implementation while still
validating the fields that make the export useful.  A workbook is first saved
beside the requested destination, reopened and verified, and only then moved
into place atomically.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
import hashlib
import math
import os
from pathlib import Path
import re
import tempfile
from typing import Any
from urllib.parse import parse_qs, unquote, urlsplit
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


__all__ = ["WorkbookStats", "validate_excel", "write_excel"]


SHEET_NAMES = ("商品汇总", "SKU明细", "图片明细", "抓取核验")
HEADER_ROW = 6
DATA_ROW = HEADER_ROW + 1
FREEZE_SPLITS = {
    "商品汇总": (5, 6),
    "SKU明细": (4, 6),
    "图片明细": (4, 6),
    "抓取核验": (0, 6),
}
TABLE_NAMES = {
    "商品汇总": "ProductsTable",
    "SKU明细": "SkuTable",
    "图片明细": "ImagesTable",
    "抓取核验": "AuditTable",
}

SUMMARY_HEADERS = (
    "全店排名",
    "列表页",
    "页内排名",
    "店铺ID",
    "商品ID",
    "商品标题",
    "当前价格 (PHP)",
    "月销原文",
    "月销下界",
    "月销状态",
    "SKU数",
    "主图链接",
    "副图数",
    "商品链接",
)
SKU_HEADERS = (
    "全店排名",
    "列表页",
    "店铺ID",
    "商品ID",
    "商品标题",
    "SKU Model ID",
    "SKU规格",
    "SKU图片链接",
    "商品链接",
)
IMAGE_HEADERS = (
    "全店排名",
    "列表页",
    "店铺ID",
    "商品ID",
    "商品标题",
    "图片类型",
    "图片序号",
    "图片链接",
    "商品链接",
)
AUDIT_HEADERS = ("核验项", "结果", "说明")
KNOWN_DETAILS_MODE = "known_product_details"
FULL_TOPSALES_MODE = "full_topsales"
KNOWN_SUMMARY_HEADERS = (
    "采集序号", "列表页（未采集）", "页内排名（未采集）",
    *SUMMARY_HEADERS[3:6], "列表价格 (PHP，未采集)", *SUMMARY_HEADERS[7:],
)
KNOWN_SKU_HEADERS = ("采集序号", "列表页（未采集）", *SKU_HEADERS[2:])
KNOWN_IMAGE_HEADERS = ("采集序号", "列表页（未采集）", *IMAGE_HEADERS[2:])
KNOWN_SUMMARY_TITLE = "UGREEN Shopee 已知商品详情刷新"
KNOWN_AUDIT_TITLE = "UGREEN Shopee 已知商品详情刷新核验"
KNOWN_SCOPE = "用户确认无新增商品的参考清单全部详情；本次未重抓列表"
KNOWN_ORDER = "参考清单采集顺序（非今日 Top Sales 排名）"
KNOWN_SKU_NOTE = "本表列出本次详情页公开 SKU；本次未采集列表价格。"
KNOWN_CONCLUSION = "参考清单全部商品详情已刷新；不代表今日 Top Sales 列表已刷新"

STORE_URL = "https://shopee.ph/ugreen.ph?page=0&shop=64922227&sortBy=sales&tab=0"
TEXT_FORMAT = "@"
CURRENCY_FORMAT = '"₱"#,##0.00'
INTEGER_FORMAT = "General"

_CORAL = "FFC9412B"
_CORAL_MEDIUM = "FFE8785C"
_CORAL_LIGHT = "FFF0A38F"
_AUDIT_GREEN = "FF7A9E7E"
_WHITE = "FFFFFFFF"
_TITLE_TEXT = "FF222222"
_META_TEXT = "FF666666"
_TEXT = "FF242424"
_LINK = "FF0563C1"
_PASS = "FF247A3D"
_FAIL = "FFB3261E"
_PASS_FILL = "FFFCE8E2"
_FAIL_FILL = "FFFCE8E6"

# TableStyleMedium2 resolves its band colours through the workbook theme.  The
# reference workbook uses this compact ChatGPT theme rather than openpyxl's
# legacy Office 2007 default, so keep it embedded to make rendering stable.
_REFERENCE_THEME = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="ChatGPT">
  <a:themeElements>
    <a:clrScheme name="ChatGPT">
      <a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>
      <a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>
      <a:dk2><a:srgbClr val="0E2841"/></a:dk2>
      <a:lt2><a:srgbClr val="E8E8E8"/></a:lt2>
      <a:accent1><a:srgbClr val="156082"/></a:accent1>
      <a:accent2><a:srgbClr val="E97132"/></a:accent2>
      <a:accent3><a:srgbClr val="196B24"/></a:accent3>
      <a:accent4><a:srgbClr val="0F9ED5"/></a:accent4>
      <a:accent5><a:srgbClr val="A02B93"/></a:accent5>
      <a:accent6><a:srgbClr val="4EA72E"/></a:accent6>
      <a:hlink><a:srgbClr val="467886"/></a:hlink>
      <a:folHlink><a:srgbClr val="96607D"/></a:folHlink>
    </a:clrScheme>
    <a:fontScheme name="Office">
      <a:majorFont>
        <a:latin typeface="Calibri Light"/><a:ea typeface="Calibri Light"/>
        <a:cs typeface="Calibri Light"/>
      </a:majorFont>
      <a:minorFont>
        <a:latin typeface="Calibri"/><a:ea typeface="Calibri"/><a:cs typeface="Calibri"/>
      </a:minorFont>
    </a:fontScheme>
    <a:fmtScheme name="ChatGPT">
      <a:fillStyleLst>
        <a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
        <a:gradFill><a:gsLst>
          <a:gs pos="0"><a:schemeClr val="phClr"><a:tint val="67000"/><a:lumMod val="110000"/><a:satMod val="105000"/></a:schemeClr></a:gs>
          <a:gs pos="50000"><a:schemeClr val="phClr"><a:tint val="73000"/><a:lumMod val="105000"/><a:satMod val="103000"/></a:schemeClr></a:gs>
          <a:gs pos="100000"><a:schemeClr val="phClr"><a:tint val="81000"/><a:lumMod val="105000"/><a:satMod val="109000"/></a:schemeClr></a:gs>
        </a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>
        <a:gradFill><a:gsLst>
          <a:gs pos="0"><a:schemeClr val="phClr"><a:tint val="94000"/><a:lumMod val="102000"/><a:satMod val="103000"/></a:schemeClr></a:gs>
          <a:gs pos="50000"><a:schemeClr val="phClr"><a:shade val="100000"/><a:lumMod val="100000"/><a:satMod val="110000"/></a:schemeClr></a:gs>
          <a:gs pos="100000"><a:schemeClr val="phClr"><a:shade val="78000"/><a:lumMod val="99000"/><a:satMod val="120000"/></a:schemeClr></a:gs>
        </a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>
      </a:fillStyleLst>
      <a:lnStyleLst>
        <a:ln w="12700"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
        <a:ln w="19050"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
        <a:ln w="25400"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
      </a:lnStyleLst>
      <a:effectStyleLst>
        <a:effectStyle><a:effectLst/></a:effectStyle>
        <a:effectStyle><a:effectLst/></a:effectStyle>
        <a:effectStyle><a:effectLst><a:outerShdw blurRad="57150" dist="19050" dir="5400000"><a:srgbClr val="000000"><a:alpha val="63000"/></a:srgbClr></a:outerShdw></a:effectLst></a:effectStyle>
      </a:effectStyleLst>
      <a:bgFillStyleLst>
        <a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
        <a:solidFill><a:schemeClr val="phClr"><a:tint val="95000"/><a:satMod val="170000"/></a:schemeClr></a:solidFill>
        <a:gradFill><a:gsLst>
          <a:gs pos="0"><a:schemeClr val="phClr"><a:tint val="93000"/><a:shade val="98000"/><a:lumMod val="102000"/><a:satMod val="150000"/></a:schemeClr></a:gs>
          <a:gs pos="50000"><a:schemeClr val="phClr"><a:tint val="98000"/><a:shade val="90000"/><a:lumMod val="103000"/><a:satMod val="130000"/></a:schemeClr></a:gs>
          <a:gs pos="100000"><a:schemeClr val="phClr"><a:shade val="63000"/><a:satMod val="120000"/></a:schemeClr></a:gs>
        </a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>
      </a:bgFillStyleLst>
    </a:fmtScheme>
  </a:themeElements>
  <a:objectDefaults/>
</a:theme>
"""

_PDP_IDENTITY_PATTERNS = (
    re.compile(r"(?:-i\.|/i\.)(\d+)\.(\d+)(?:/)?$", re.I),
    re.compile(r"/product/(\d+)/(\d+)(?:/)?$", re.I),
)

@dataclass(frozen=True, slots=True)
class WorkbookStats:
    """Counts proven by reopening the saved workbook."""

    product_count: int
    sku_count: int
    main_image_count: int
    secondary_image_count: int
    audit_status: str
    collection_mode: str = FULL_TOPSALES_MODE


@dataclass(frozen=True, slots=True)
class _SkuView:
    model_id: str
    props: str
    image_url: str | None


@dataclass(frozen=True, slots=True)
class _ProductView:
    global_rank: int
    list_page: int | None
    source_position: int | None
    shop_id: str
    item_id: str
    title: str
    product_url: str
    price_php: int | float | None
    monthly_sales_text: str | None
    monthly_sales_lower_bound: int | None
    monthly_sales_status: str
    skus: tuple[_SkuView, ...]
    main_image_url: str | None
    secondary_image_urls: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class _ExportView:
    store_url: str
    captured_at: str
    products: tuple[_ProductView, ...]
    audit: Any
    collection_mode: str
    scope_metadata: Mapping[str, Any]


@dataclass(frozen=True, slots=True)
class _Expected:
    product_count: int
    sku_count: int
    main_image_count: int
    secondary_image_count: int
    summary_hyperlink_count: int
    sku_hyperlink_count: int
    image_hyperlink_count: int
    audit_row_count: int
    audit_status: str
    collection_mode: str = FULL_TOPSALES_MODE


def _has(obj: Any, name: str) -> bool:
    if isinstance(obj, Mapping):
        return name in obj
    return hasattr(obj, name)


def _get(obj: Any, *names: str, default: Any = None) -> Any:
    saw_none = False
    for name in names:
        if isinstance(obj, Mapping):
            if name not in obj:
                continue
            value = obj[name]
        elif hasattr(obj, name):
            value = getattr(obj, name)
        else:
            continue
        if value is not None:
            return value
        saw_none = True
    return None if saw_none else default


def _as_sequence(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, (str, bytes, bytearray)):
        return [value]
    if isinstance(value, Mapping):
        return list(value.values())
    if isinstance(value, Sequence):
        return list(value)
    try:
        return list(value)
    except TypeError:
        return [value]


def _clean_text(value: Any, *, limit: int = 32_767) -> str:
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    return text[:limit]


def _id_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _clean_text(value)


def _integer(value: Any, *, default: int) -> int:
    if value is None or value == "":
        return default
    try:
        return int(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError(f"应为整数，实际为 {value!r}") from exc


def _number(value: Any) -> int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"布尔值不能作为有效数值：{value!r}")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError(f"数值不是有限值：{value!r}")
        return int(value) if value.is_integer() else value
    try:
        decimal_value = value if isinstance(value, Decimal) else Decimal(str(value).replace(",", ""))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"应为数值，实际为 {value!r}") from exc
    if not decimal_value.is_finite():
        raise ValueError(f"数值不是有限值：{value!r}")
    if decimal_value == decimal_value.to_integral_value():
        return int(decimal_value)
    return float(decimal_value)


def _optional_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    return _integer(value, default=0)


def _format_props(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Mapping):
        parts: list[str] = []
        for key, item in value.items():
            if item is None or item == "":
                continue
            if isinstance(item, (list, tuple)):
                item_text = " / ".join(_clean_text(part) for part in item)
            else:
                item_text = _clean_text(item)
            parts.append(f"{_clean_text(key)}: {item_text}")
        return ", ".join(parts)
    if isinstance(value, (list, tuple)):
        return ", ".join(_clean_text(item) for item in value)
    return _clean_text(value)


def _monthly_status(observation: Any, sales_text: str | None) -> str:
    if observation is None:
        return "已展示" if sales_text else "未展示"
    normalized = str(observation).strip().lower().replace("-", "_").replace(" ", "_")
    if normalized in {"not_collected", "未采集"}:
        return "未采集"
    if normalized in {"displayed", "shown", "observed", "已展示"}:
        return "已展示"
    if normalized in {"not_displayed", "not_shown", "missing", "unobserved", "未展示"}:
        return "未展示"
    if "not" in normalized or "未" in normalized or "missing" in normalized:
        return "未展示"
    return "已展示" if sales_text else "未展示"


def _local_excel_datetime(value: Any) -> str:
    """Return the same local, display-ready timestamp used by the gold workbook."""

    if value is None:
        timestamp = datetime.now().astimezone()
    elif isinstance(value, datetime):
        timestamp = value.astimezone() if value.tzinfo is not None else value
    else:
        text = _clean_text(value).strip()
        if re.fullmatch(r"\d{4}/\d{2}/\d{2} \d{2}:\d{2}:\d{2}", text):
            return text
        try:
            timestamp = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError(f"采集时间无法解析：{text!r}") from exc
        if timestamp.tzinfo is not None:
            timestamp = timestamp.astimezone()
    return timestamp.strftime("%Y/%m/%d %H:%M:%S")


def _normalise_export(result: Any) -> _ExportView:
    collection_mode = _get(result, "collection_mode", default=FULL_TOPSALES_MODE)
    if collection_mode not in {FULL_TOPSALES_MODE, KNOWN_DETAILS_MODE}:
        raise ValueError("未知的工作簿采集模式")
    known_details = collection_mode == KNOWN_DETAILS_MODE
    raw_products = _as_sequence(_get(result, "products", default=[]))
    if not raw_products:
        raise ValueError("没有商品数据，无法创建 Excel 工作簿")

    products: list[_ProductView] = []
    for index, raw in enumerate(raw_products, start=1):
        if known_details:
            for field in ("global_rank", "rank", "source_page", "list_page", "page",
                          "source_position", "page_rank", "position", "price_php",
                          "current_price", "price", "monthly_sales_text", "monthly_sales_raw",
                          "monthly_sales_count_lower_bound", "monthly_sales_lower_bound"):
                if _get(raw, field) is not None:
                    raise ValueError(f"已知商品详情刷新不得携带旧列表字段：{field}")
            if _get(raw, "monthly_sales_observation", "monthly_sales_status") != "not_collected":
                raise ValueError("已知商品详情刷新的月销状态必须为 not_collected")
            global_rank, list_page, source_position = index, None, None
        else:
            global_rank = _integer(_get(raw, "global_rank", "rank"), default=index)
            if _has(raw, "source_page"):
                list_page = _integer(_get(raw, "source_page"), default=0) + 1
            else:
                list_page = _integer(_get(raw, "list_page", "page"), default=1)
            source_position = _integer(
                _get(raw, "source_position", "page_rank", "position"),
                default=((global_rank - 1) % 30) + 1,
            )
        shop_id = _id_text(_get(raw, "shop_id"))
        item_id = _id_text(_get(raw, "item_id"))
        title = _clean_text(_get(raw, "title", "name", default=""))
        product_url = _clean_text(_get(raw, "product_url", "url", default=""))
        if not shop_id or not item_id:
            raise ValueError(f"排名为 {global_rank} 的商品缺少 shop_id 或 item_id")
        if not title:
            raise ValueError(f"商品 {shop_id}.{item_id} 缺少标题")
        if not _is_http_url(product_url):
            raise ValueError(f"商品 {shop_id}.{item_id} 的 URL 无效：{product_url!r}")

        sales_text_value = _get(raw, "monthly_sales_text", "monthly_sales_raw")
        sales_text = _clean_text(sales_text_value) if sales_text_value not in (None, "") else None
        lower_bound = _optional_int(
            _get(raw, "monthly_sales_count_lower_bound", "monthly_sales_lower_bound")
        )
        status = _monthly_status(_get(raw, "monthly_sales_observation", "monthly_sales_status"), sales_text)
        if not known_details and status == "未采集":
            raise ValueError("完整 Top Sales 模式不能使用未采集月销状态")

        skus: list[_SkuView] = []
        for sku in _as_sequence(_get(raw, "skus", default=[])):
            skus.append(
                _SkuView(
                    model_id=_id_text(_get(sku, "model_id", "sku_model_id", "id")),
                    props=_format_props(_get(sku, "sku_props", "props", "specification")),
                    image_url=_optional_text(_get(sku, "sku_image_url", "image_url")),
                )
            )

        main_image_url = _optional_text(_get(raw, "main_image_url", "main_image"))
        secondary_images = tuple(
            text
            for value in _as_sequence(
                _get(raw, "secondary_image_urls", "secondary_images", default=[])
            )
            if (text := _optional_text(value)) is not None
        )
        products.append(
            _ProductView(
                global_rank=global_rank,
                list_page=list_page,
                source_position=source_position,
                shop_id=shop_id,
                item_id=item_id,
                title=title,
                product_url=product_url,
                price_php=_number(_get(raw, "price_php", "current_price", "price")),
                monthly_sales_text=sales_text,
                monthly_sales_lower_bound=lower_bound,
                monthly_sales_status=status,
                skus=tuple(skus),
                main_image_url=main_image_url,
                secondary_image_urls=secondary_images,
            )
        )

    ranks = [product.global_rank for product in products]
    if ranks != list(range(1, len(products) + 1)):
        raise ValueError("商品 global_rank 值必须从 1 开始连续排列")
    identities = [(product.shop_id, product.item_id) for product in products]
    if len(set(identities)) != len(identities):
        raise ValueError("商品列表中仍存在重复的 shop_id + item_id 标识")
    scope_metadata = _get(result, "scope_metadata", default={})
    if not isinstance(scope_metadata, Mapping):
        raise ValueError("采集范围元数据必须为映射")
    if known_details:
        _validate_known_scope(scope_metadata, identities)

    source_store_url = _clean_text(_get(result, "store_url", default=STORE_URL))
    if not _is_http_url(source_store_url):
        raise ValueError(f"店铺 URL 无效：{source_store_url!r}")
    return _ExportView(
        # The crawler may omit the redundant ``shop`` query parameter.  The
        # exported workbook deliberately uses the stable canonical display
        # URL from the verified reference workbook.
        store_url=STORE_URL,
        captured_at=_local_excel_datetime(_get(result, "captured_at", "scraped_at")),
        products=tuple(products),
        audit=_get(result, "audit"),
        collection_mode=collection_mode,
        scope_metadata=scope_metadata,
    )


def _identity_sha256(identities: Sequence[tuple[str, str]]) -> str:
    return hashlib.sha256("\n".join(f"{shop}:{item}" for shop, item in identities).encode("utf-8")).hexdigest()


def _validate_known_scope(
    metadata: Mapping[str, Any], identities: Sequence[tuple[str, str]], *, reference_name_only: bool = False,
) -> None:
    reference = metadata.get("reference_workbook")
    if not isinstance(reference, str) or not reference or Path(reference).suffix.lower() != ".xlsx":
        raise ValueError("已知商品详情刷新缺少有效参考工作簿")
    if reference_name_only:
        if Path(reference).name != reference or "\\" in reference:
            raise ValueError("参考工作簿审计只应保留文件名")
    elif not Path(reference).is_absolute():
        raise ValueError("已知商品详情刷新缺少参考工作簿绝对路径")
    for key in ("reference_sha256", "reference_identity_sha256"):
        if not isinstance(metadata.get(key), str) or re.fullmatch(r"[0-9a-f]{64}", metadata[key]) is None:
            raise ValueError(f"已知商品详情刷新的 {key} 无效")
    captured = metadata.get("reference_captured_at")
    try:
        if not isinstance(captured, str) or re.fullmatch(r"\d{4}/\d{2}/\d{2} \d{2}:\d{2}:\d{2}", captured) is None:
            raise ValueError
        reference_date = datetime.strptime(captured, "%Y/%m/%d %H:%M:%S").date().isoformat()
    except ValueError as exc:
        raise ValueError("参考工作簿采集时间无效") from exc
    if metadata.get("reference_date") != reference_date:
        raise ValueError("参考工作簿日期与采集时间不一致")
    if type(metadata.get("reference_unique_count")) is not int or metadata["reference_unique_count"] != len(identities):
        raise ValueError("参考唯一商品数与本次详情数量不一致")
    if metadata["reference_identity_sha256"] != _identity_sha256(identities):
        raise ValueError("参考商品身份序列与本次详情不一致")
    if metadata.get("user_confirmed_no_new_products") is not True:
        raise ValueError("缺少用户确认今日无新增商品")


def _optional_text(value: Any) -> str | None:
    if value is None or value == "":
        return None
    return _clean_text(value)


def _is_http_url(value: str | None) -> bool:
    if not value:
        return False
    parsed = urlsplit(value)
    return parsed.scheme.lower() in {"http", "https"} and bool(parsed.netloc)


def _audit_int(audit: Any, name: str, fallback: int) -> int:
    return _integer(_get(audit, name), default=fallback)


def _audit_rows(view: _ExportView) -> tuple[list[tuple[Any, Any, str]], str]:
    products = view.products
    product_count = len(products)
    sku_count = sum(len(product.skus) for product in products)
    main_count = sum(bool(product.main_image_url) for product in products)
    secondary_count = sum(len(product.secondary_image_urls) for product in products)
    products_with_skus = sum(bool(product.skus) for product in products)
    missing_monthly = sum(product.monthly_sales_status == "未展示" for product in products)
    known_details = view.collection_mode == KNOWN_DETAILS_MODE

    page_count = _audit_int(view.audit, "list_page_count", 0 if known_details else max(product.list_page for product in products))
    occurrences = _audit_int(view.audit, "list_input_product_occurrences", product_count)
    duplicate_count = _audit_int(
        view.audit,
        "list_duplicate_occurrence_count",
        max(0, occurrences - product_count),
    )
    detail_success = _audit_int(view.audit, "detail_success_count", product_count)
    detail_failure = _audit_int(view.audit, "detail_failure_count", 0)
    detail_parse_success = _audit_int(
        view.audit,
        "detail_parse_success_count",
        detail_success,
    )
    detail_parse_errors = _audit_int(
        view.audit,
        "detail_parse_error_count",
        detail_failure,
    )
    listed_product_count = _audit_int(view.audit, "listed_product_count", product_count)
    audit_missing_monthly = _audit_int(
        view.audit,
        "products_without_monthly_sales_display",
        missing_monthly,
    )

    invalid_product_urls = sum(not _is_http_url(product.product_url) for product in products)
    invalid_image_urls = 0
    for product in products:
        invalid_image_urls += not _is_http_url(product.main_image_url)
        invalid_image_urls += sum(not _is_http_url(url) for url in product.secondary_image_urls)
        invalid_image_urls += sum(
            bool(sku.image_url) and not _is_http_url(sku.image_url) for sku in product.skus
        )
    ranks_contiguous = [product.global_rank for product in products] == list(range(1, product_count + 1))

    list_scope_passed = (
        page_count == occurrences == duplicate_count == 0
        and all(product.monthly_sales_status == "未采集" for product in products)
        if known_details else
        page_count > 0 and occurrences - duplicate_count == product_count
    )
    if known_details:
        required_counts = {
            "list_page_count": 0, "list_input_product_occurrences": 0,
            "list_duplicate_occurrence_count": 0, "listed_product_count": product_count,
            "detail_success_count": product_count, "detail_failure_count": 0,
            "products_without_monthly_sales_display": 0,
        }
        for field, expected_count in required_counts.items():
            value = _get(view.audit, field)
            if type(value) is not int or value != expected_count:
                raise ValueError(f"已知商品详情刷新审计字段 {field} 不一致")
    passed = (
        product_count > 0
        and list_scope_passed
        and listed_product_count == product_count
        and ranks_contiguous
        and detail_success == product_count
        and detail_failure == 0
        and detail_parse_success == product_count
        and detail_parse_errors == 0
        and audit_missing_monthly == missing_monthly
        and products_with_skus == product_count
        and main_count == product_count
        and invalid_product_urls == 0
        and invalid_image_urls == 0
    )
    status = "通过" if passed else "未通过"
    rows: list[tuple[Any, Any, str]] = [
        ("店铺 URL", view.store_url, "固定抓取 UGREEN 菲律宾官方店铺"),
        ("排序方式", "Top Sales", "列表 URL 使用 sortBy=sales"),
        ("列表总页数", page_count, "动态遍历所有分页"),
        ("商品卡出现次数", occurrences, "包含跨页重复展示"),
        ("跨页重复次数", duplicate_count, "按 shop_id + item_id 去重"),
        ("唯一商品数", product_count, "商品汇总行数"),
        ("成功进入详情页", detail_success, "浏览器真实导航并匹配 PDP 身份"),
        ("详情页失败数", detail_failure, "重试后仍失败的商品"),
        ("详情解析成功数", detail_parse_success, "严格匹配详情数据"),
        ("详情解析错误数", detail_parse_errors, "应为 0"),
        ("未展示月销商品数", audit_missing_monthly, "留空/标记未展示，不按 0 处理"),
        ("SKU 总数", sku_count, "SKU 明细行数"),
        ("有 SKU 的商品数", products_with_skus, ""),
        ("主图数", main_count, ""),
        ("副图数", secondary_count, ""),
        ("无效图片 URL 数", invalid_image_urls, "应为 0"),
        (
            "完整性结论",
            status,
            "全部唯一商品已进入详情页并写入 Excel" if passed else "存在未通过的完整性检查",
        ),
    ]
    if known_details:
        metadata = view.scope_metadata
        rows = [
            ("店铺 URL", view.store_url, "固定范围为 UGREEN 菲律宾官方店铺"),
            ("采集模式", KNOWN_DETAILS_MODE, "已知商品详情刷新，不是本次列表全量刷新"),
            ("参考工作簿", Path(metadata["reference_workbook"]).name, "只读取商品身份和链接，不复用旧业务数据"),
            ("参考日期", metadata["reference_date"], "参考清单原采集日期"),
            ("参考采集时间", metadata["reference_captured_at"], "参考工作簿商品汇总 B3 原文"),
            ("参考文件 SHA-256", metadata["reference_sha256"], "参考工作簿原始文件摘要"),
            ("参考身份序列 SHA-256", metadata["reference_identity_sha256"], "按参考顺序，以换行拼接 shop_id:item_id 后计算"),
            ("参考唯一商品数", metadata["reference_unique_count"], "本次必须逐个实际访问全部参考商品"),
            ("用户确认无新增商品", "是", "范围依据用户确认，不代表本次独立发现全部商品"),
            ("本次重抓列表", "否", "列表排名、页码、价格和月销均未采集"),
            ("排序方式", KNOWN_ORDER, "采集序号仅为访问顺序"),
            ("列表总页数", 0, "本次未采集列表，非店铺没有列表"),
            ("商品卡出现次数", 0, "本次未采集列表"),
            ("跨页重复次数", 0, "本次未采集列表"),
            *rows[5:10],
            ("未采集月销商品数", product_count, "未采集不等于未展示，更不等于零"),
            *rows[11:16],
            ("完整性结论", status, KNOWN_CONCLUSION if passed else "参考清单详情完整性检查未通过"),
        ]
    return rows, status


def _set_text(cell: Cell, value: Any, *, identifier: bool = False) -> None:
    if value is None or value == "":
        cell.value = None
    else:
        cell.value = _clean_text(value)
        # openpyxl interprets leading '=' as a formula unless its type is reset.
        cell.data_type = "s"
    if identifier:
        cell.number_format = TEXT_FORMAT


def _set_number(cell: Cell, value: int | float | None, number_format: str) -> None:
    cell.value = value
    cell.number_format = number_format


def _set_link(cell: Cell, value: str | None) -> None:
    _set_text(cell, value)
    if value and _is_http_url(value):
        cell.hyperlink = value
        cell.font = Font(name="Arial", size=10, color=_LINK)


def _style_title(ws: Any, title: str, last_column: int) -> None:
    for column in range(1, last_column + 1):
        cell = ws.cell(row=2, column=column)
        cell.font = Font(name="Arial", size=14, bold=True, color=_TITLE_TEXT)
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border(bottom=Side(style="thin", color=_CORAL))
        cell.alignment = Alignment(vertical="center")
    cell = ws.cell(row=2, column=1)
    _set_text(cell, title)
    ws.row_dimensions[2].height = 26


def _style_context_rows(ws: Any, last_column: int) -> None:
    for row in (3, 4):
        for column in range(1, last_column + 1):
            cell = ws.cell(row=row, column=column)
            cell.font = Font(name="Arial", size=10, color=_META_TEXT)
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 21
    for column in range(1, last_column + 1):
        cell = ws.cell(row=5, column=column)
        cell.font = Font(name="Arial", size=10, color=_TEXT)
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()
        cell.alignment = Alignment(vertical="center")


def _style_metadata_pair(ws: Any, row: int, label_column: int, label: str, value: Any) -> None:
    label_cell = ws.cell(row=row, column=label_column)
    value_cell = ws.cell(row=row, column=label_column + 1)
    _set_text(label_cell, label)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        value_cell.value = value
        value_cell.number_format = INTEGER_FORMAT
    else:
        _set_text(value_cell, value)


def _style_headers(ws: Any, headers: Sequence[str]) -> None:
    separator = Side(style="thin", color=_WHITE)
    underline = Side(style="medium", color=_CORAL)
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=column)
        _set_text(cell, header)
        cell.font = Font(name="Arial", size=10, bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_CORAL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            left=separator if column > 1 else Side(),
            right=separator if column < len(headers) else Side(),
            bottom=underline,
        )
    ws.row_dimensions[HEADER_ROW].height = 30


def _configure_sheet(ws: Any, *, tab_color: str, freeze_panes: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze_panes
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_format.defaultRowHeight = 16.8
    ws.sheet_format.baseColWidth = 8
    ws.auto_filter.ref = None


def _set_widths(ws: Any, widths: Mapping[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _add_table(ws: Any, *, name: str, last_row: int, last_column: int) -> None:
    reference = f"A{HEADER_ROW}:{get_column_letter(last_column)}{max(last_row, HEADER_ROW)}"
    if last_row <= HEADER_ROW:
        return
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _style_data_cell(
    cell: Cell,
    *,
    horizontal: str | None = None,
    vertical: str = "center",
    wrap: bool = False,
) -> None:
    if cell.hyperlink is None:
        cell.font = Font(name="Arial", size=10, color=_TEXT)
    cell.alignment = Alignment(
        horizontal=horizontal,
        vertical=vertical,
        wrap_text=wrap,
    )
    cell.border = Border()


def _build_summary_sheet(wb: Workbook, view: _ExportView) -> None:
    ws = wb.active
    ws.title = SHEET_NAMES[0]
    _configure_sheet(ws, tab_color=_CORAL, freeze_panes="F7")
    known_details = view.collection_mode == KNOWN_DETAILS_MODE
    _style_title(ws, KNOWN_SUMMARY_TITLE if known_details else "UGREEN Shopee Top Sales 全店商品", len(SUMMARY_HEADERS))
    _style_context_rows(ws, len(SUMMARY_HEADERS))

    product_count = len(view.products)
    sku_count = sum(len(product.skus) for product in view.products)
    main_count = sum(bool(product.main_image_url) for product in view.products)
    secondary_count = sum(len(product.secondary_image_urls) for product in view.products)
    _style_metadata_pair(ws, 3, 1, "采集时间", view.captured_at)
    _style_metadata_pair(ws, 3, 5, "商品数", product_count)
    _style_metadata_pair(ws, 3, 7, "SKU数", sku_count)
    _style_metadata_pair(ws, 3, 9, "主图数", main_count)
    _style_metadata_pair(ws, 3, 11, "副图数", secondary_count)
    _style_metadata_pair(ws, 4, 1, "店铺来源", view.store_url)
    _style_metadata_pair(ws, 4, 5, "排序", KNOWN_ORDER if known_details else "Top Sales")
    _style_headers(ws, KNOWN_SUMMARY_HEADERS if known_details else SUMMARY_HEADERS)

    for row_number, product in enumerate(view.products, start=DATA_ROW):
        values = (
            product.global_rank,
            product.list_page,
            product.source_position,
            product.shop_id,
            product.item_id,
            product.title,
            product.price_php,
            product.monthly_sales_text,
            product.monthly_sales_lower_bound,
            product.monthly_sales_status,
            len(product.skus),
            product.main_image_url,
            len(product.secondary_image_urls),
            product.product_url,
        )
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_number, column=column)
            if column in {4, 5}:
                _set_text(cell, value, identifier=True)
            elif column in {12, 14}:
                _set_link(cell, value)
            elif column == 7:
                _set_number(cell, value, CURRENCY_FORMAT)
            elif column in {1, 2, 3, 9, 11, 13}:
                _set_number(cell, value, INTEGER_FORMAT)
            else:
                _set_text(cell, value)
            if column == 6:
                _style_data_cell(cell, vertical="top", wrap=True)
            elif column in {7, 11, 13}:
                _style_data_cell(cell, horizontal="right", wrap=column == 13)
            elif column in {12, 14}:
                _style_data_cell(cell, wrap=True)
            else:
                _style_data_cell(cell, horizontal="center")
        ws.row_dimensions[row_number].height = 42

    _set_widths(
        ws,
        {
            "A": 9,
            "D": 15,
            "F": 46,
            "G": 17,
            "H": 19,
            "I": 13,
            "K": 10,
            "L": 42,
            "M": 10,
            "N": 46,
        },
    )
    _add_table(
        ws,
        name="ProductsTable",
        last_row=HEADER_ROW + product_count,
        last_column=len(SUMMARY_HEADERS),
    )


def _build_sku_sheet(wb: Workbook, view: _ExportView) -> None:
    ws = wb.create_sheet(SHEET_NAMES[1])
    _configure_sheet(ws, tab_color=_CORAL_MEDIUM, freeze_panes="E7")
    _style_title(ws, "UGREEN Shopee SKU 明细", len(SKU_HEADERS))
    _style_context_rows(ws, len(SKU_HEADERS))
    sku_count = sum(len(product.skus) for product in view.products)
    _style_metadata_pair(ws, 3, 1, "记录数", sku_count)
    _style_metadata_pair(ws, 3, 5, "商品数", len(view.products))
    known_details = view.collection_mode == KNOWN_DETAILS_MODE
    _style_metadata_pair(ws, 4, 1, "说明", KNOWN_SKU_NOTE if known_details else "本表列出详情页公开 SKU；商品当前价格见“商品汇总”。")
    _style_headers(ws, KNOWN_SKU_HEADERS if known_details else SKU_HEADERS)

    row_number = DATA_ROW
    for product in view.products:
        for sku in product.skus:
            values = (
                product.global_rank,
                product.list_page,
                product.shop_id,
                product.item_id,
                product.title,
                sku.model_id,
                sku.props,
                sku.image_url,
                product.product_url,
            )
            for column, value in enumerate(values, start=1):
                cell = ws.cell(row=row_number, column=column)
                if column in {3, 4, 6}:
                    _set_text(cell, value, identifier=True)
                elif column in {8, 9}:
                    _set_link(cell, value)
                elif column in {1, 2}:
                    _set_number(cell, value, INTEGER_FORMAT)
                else:
                    _set_text(cell, value)
                _style_data_cell(
                    cell,
                    horizontal="center" if column in {1, 2, 3, 4, 6} else None,
                    wrap=column in {5, 6, 7, 8, 9},
                )
            row_number += 1
            ws.row_dimensions[row_number - 1].height = 40

    _set_widths(
        ws,
        {"A": 9, "C": 15, "E": 44, "F": 18, "G": 34, "H": 42, "I": 46},
    )
    _add_table(
        ws,
        name="SkuTable",
        last_row=row_number - 1,
        last_column=len(SKU_HEADERS),
    )


def _build_image_sheet(wb: Workbook, view: _ExportView) -> None:
    ws = wb.create_sheet(SHEET_NAMES[2])
    _configure_sheet(ws, tab_color=_CORAL_LIGHT, freeze_panes="E7")
    _style_title(ws, "UGREEN Shopee 商品图片明细", len(IMAGE_HEADERS))
    _style_context_rows(ws, len(IMAGE_HEADERS))
    main_count = sum(bool(product.main_image_url) for product in view.products)
    secondary_count = sum(len(product.secondary_image_urls) for product in view.products)
    _style_metadata_pair(ws, 3, 1, "图片记录数", main_count + secondary_count)
    _style_metadata_pair(ws, 3, 5, "主图数", main_count)
    _style_metadata_pair(ws, 3, 7, "副图数", secondary_count)
    _style_metadata_pair(ws, 4, 1, "说明", "每个商品的主图列在前，副图保持详情页图库顺序。")
    _style_headers(ws, KNOWN_IMAGE_HEADERS if view.collection_mode == KNOWN_DETAILS_MODE else IMAGE_HEADERS)

    row_number = DATA_ROW
    for product in view.products:
        image_rows: list[tuple[str, int, str]] = []
        if product.main_image_url:
            image_rows.append(("主图", 1, product.main_image_url))
        image_rows.extend(
            ("副图", image_index, image_url)
            for image_index, image_url in enumerate(product.secondary_image_urls, start=1)
        )
        for image_type, image_index, image_url in image_rows:
            values = (
                product.global_rank,
                product.list_page,
                product.shop_id,
                product.item_id,
                product.title,
                image_type,
                image_index,
                image_url,
                product.product_url,
            )
            for column, value in enumerate(values, start=1):
                cell = ws.cell(row=row_number, column=column)
                if column in {3, 4}:
                    _set_text(cell, value, identifier=True)
                elif column in {8, 9}:
                    _set_link(cell, value)
                elif column in {1, 2, 7}:
                    _set_number(cell, value, INTEGER_FORMAT)
                else:
                    _set_text(cell, value)
                _style_data_cell(
                    cell,
                    horizontal="center" if column in {1, 2, 3, 4, 6, 7} else None,
                    wrap=column in {5, 8, 9},
                )
            row_number += 1
            ws.row_dimensions[row_number - 1].height = 40

    _set_widths(
        ws,
        {"A": 9, "C": 15, "E": 44, "F": 11, "H": 50, "I": 46},
    )
    _add_table(
        ws,
        name="ImagesTable",
        last_row=row_number - 1,
        last_column=len(IMAGE_HEADERS),
    )


def _build_audit_sheet(
    wb: Workbook,
    view: _ExportView,
    rows: Sequence[tuple[Any, Any, str]],
    status: str,
) -> None:
    ws = wb.create_sheet(SHEET_NAMES[3])
    _configure_sheet(
        ws,
        tab_color=_AUDIT_GREEN if status == "通过" else _FAIL,
        freeze_panes="A7",
    )
    known_details = view.collection_mode == KNOWN_DETAILS_MODE
    _style_title(ws, KNOWN_AUDIT_TITLE if known_details else "UGREEN Shopee 全店抓取核验", len(AUDIT_HEADERS))
    _style_context_rows(ws, len(AUDIT_HEADERS))
    _style_metadata_pair(ws, 3, 1, "核验状态", status)
    status_cell = ws.cell(row=3, column=2)
    status_cell.font = Font(
        name="Arial",
        size=10,
        bold=True,
        color=_PASS if status == "通过" else _FAIL,
    )
    status_cell.fill = PatternFill(
        "solid",
        fgColor=_PASS_FILL if status == "通过" else _FAIL_FILL,
    )
    status_cell.alignment = Alignment(horizontal="center", vertical="center")
    _style_metadata_pair(ws, 4, 1, "范围", KNOWN_SCOPE if known_details else "UGREEN Top Sales 全部分页与全部详情页")
    _style_headers(ws, AUDIT_HEADERS)

    for row_number, values in enumerate(rows, start=DATA_ROW):
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_number, column=column)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                _set_number(cell, value, INTEGER_FORMAT)
            else:
                _set_text(cell, value)
            _style_data_cell(
                cell,
                horizontal="center" if column == 2 else None,
                wrap=column in {2, 3},
            )
            if column == 1:
                cell.font = Font(name="Arial", size=10, bold=True, color=_META_TEXT)
        if values[0] == "完整性结论":
            result_cell = ws.cell(row=row_number, column=2)
            result_cell.font = Font(
                name="Arial",
                size=10,
                bold=True,
                color=_PASS if status == "通过" else _FAIL,
            )
            result_cell.fill = PatternFill(
                "solid",
                fgColor=_PASS_FILL if status == "通过" else _FAIL_FILL,
            )
        ws.row_dimensions[row_number].height = 46 if row_number == DATA_ROW else 28

    _set_widths(ws, {"A": 26, "B": 30, "C": 54})
    _add_table(
        ws,
        name="AuditTable",
        last_row=HEADER_ROW + len(rows),
        last_column=len(AUDIT_HEADERS),
    )


def _build_workbook(view: _ExportView) -> tuple[Workbook, _Expected]:
    audit_rows, audit_status = _audit_rows(view)
    wb = Workbook()
    wb.loaded_theme = _REFERENCE_THEME.encode("utf-8")

    _build_summary_sheet(wb, view)
    _build_sku_sheet(wb, view)
    _build_image_sheet(wb, view)
    _build_audit_sheet(wb, view, audit_rows, audit_status)
    wb.active = 0

    summary_hyperlinks = sum(bool(product.main_image_url) + 1 for product in view.products)
    sku_hyperlinks = sum(
        1 + bool(sku.image_url)
        for product in view.products
        for sku in product.skus
    )
    image_hyperlinks = 2 * sum(
        bool(product.main_image_url) + len(product.secondary_image_urls)
        for product in view.products
    )
    expected = _Expected(
        product_count=len(view.products),
        sku_count=sum(len(product.skus) for product in view.products),
        main_image_count=sum(bool(product.main_image_url) for product in view.products),
        secondary_image_count=sum(len(product.secondary_image_urls) for product in view.products),
        summary_hyperlink_count=summary_hyperlinks,
        sku_hyperlink_count=sku_hyperlinks,
        image_hyperlink_count=image_hyperlinks,
        audit_row_count=len(audit_rows),
        audit_status=audit_status,
        collection_mode=view.collection_mode,
    )
    return wb, expected


def _assert_headers(ws: Any, expected: Sequence[str]) -> None:
    if ws.max_column != len(expected):
        raise ValueError(
            f"工作表 {ws.title!r} 的列数异常："
            f"应为 {len(expected)}，实际为 {ws.max_column}"
        )
    actual = tuple(ws.cell(row=HEADER_ROW, column=column).value for column in range(1, len(expected) + 1))
    if actual != tuple(expected):
        raise ValueError(f"工作表 {ws.title!r} 的表头异常：{actual!r}")


def _assert_reference_layout(ws: Any) -> None:
    expected_x, expected_y = FREEZE_SPLITS[ws.title]
    pane = ws.sheet_view.pane
    actual_x = int(pane.xSplit or 0) if pane is not None else 0
    actual_y = int(pane.ySplit or 0) if pane is not None else 0
    if pane is None or pane.state != "frozen" or (actual_x, actual_y) != (expected_x, expected_y):
        raise ValueError(
            f"工作表 {ws.title!r} 的冻结窗格异常："
            f"应冻结 {expected_x} 列和 {expected_y} 行，实际为 {actual_x} 列和 {actual_y} 行"
        )
    if ws.sheet_view.showGridLines is not False:
        raise ValueError(f"工作表 {ws.title!r} 未隐藏网格线")
    if ws.merged_cells.ranges:
        raise ValueError(f"工作表 {ws.title!r} 不应包含合并单元格")
    tables = list(ws.tables.values())
    expected_name = TABLE_NAMES[ws.title]
    expected_ref = f"A{HEADER_ROW}:{get_column_letter(ws.max_column)}{ws.max_row}"
    if len(tables) != 1:
        raise ValueError(f"工作表 {ws.title!r} 应包含且仅包含一个 Excel 表")
    table = tables[0]
    style_name = table.tableStyleInfo.name if table.tableStyleInfo is not None else None
    if table.displayName != expected_name or table.ref != expected_ref:
        raise ValueError(
            f"工作表 {ws.title!r} 的 Excel 表异常："
            f"应为 {expected_name}({expected_ref})，实际为 {table.displayName}({table.ref})"
        )
    if style_name != "TableStyleMedium2":
        raise ValueError(f"工作表 {ws.title!r} 未使用 TableStyleMedium2")


def _assert_ids_are_text(ws: Any, rows: range, columns: Sequence[int]) -> None:
    for row in rows:
        for column in columns:
            cell = ws.cell(row=row, column=column)
            if cell.value in (None, ""):
                continue
            if cell.data_type != "s" or not isinstance(cell.value, str):
                raise ValueError(f"标识符 {ws.title}!{cell.coordinate} 未按文本存储")
            if cell.number_format != TEXT_FORMAT:
                raise ValueError(f"标识符 {ws.title}!{cell.coordinate} 缺少文本格式")


def _assert_links(ws: Any, columns: Sequence[int], expected_count: int) -> None:
    count = 0
    for row in range(DATA_ROW, ws.max_row + 1):
        for column in columns:
            cell = ws.cell(row=row, column=column)
            if cell.value in (None, ""):
                continue
            if cell.hyperlink is None or cell.hyperlink.target != cell.value:
                raise ValueError(f"{ws.title}!{cell.coordinate} 缺少超链接")
            count += 1
    if count != expected_count:
        raise ValueError(
            f"工作表 {ws.title!r} 的超链接数量不匹配：应为 {expected_count}，实际为 {count}"
        )


def _verify_zip(path: Path) -> None:
    if not zipfile.is_zipfile(path):
        raise ValueError(f"生成的文件不是有效的 XLSX ZIP 容器：{path}")
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise ValueError(f"XLSX ZIP 成员损坏：{bad_member}")
        members = set(archive.namelist())
        required = {"[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml"}
        missing = required - members
        if missing:
            raise ValueError(f"XLSX 容器缺少必需成员：{sorted(missing)}")


def _validate_with_expected(path: Path, expected: _Expected) -> WorkbookStats:
    _verify_zip(path)
    wb = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    try:
        if tuple(wb.sheetnames) != SHEET_NAMES:
            raise ValueError(f"工作表顺序异常：{wb.sheetnames!r}")

        summary = wb[SHEET_NAMES[0]]
        sku = wb[SHEET_NAMES[1]]
        images = wb[SHEET_NAMES[2]]
        audit = wb[SHEET_NAMES[3]]
        known_details = expected.collection_mode == KNOWN_DETAILS_MODE
        _assert_headers(summary, KNOWN_SUMMARY_HEADERS if known_details else SUMMARY_HEADERS)
        _assert_headers(sku, KNOWN_SKU_HEADERS if known_details else SKU_HEADERS)
        _assert_headers(images, KNOWN_IMAGE_HEADERS if known_details else IMAGE_HEADERS)
        _assert_headers(audit, AUDIT_HEADERS)

        expected_metadata = {
            SHEET_NAMES[0]: {
                "A2": "UGREEN Shopee Top Sales 全店商品",
                "A3": "采集时间",
                "E3": "商品数",
                "F3": expected.product_count,
                "G3": "SKU数",
                "H3": expected.sku_count,
                "I3": "主图数",
                "J3": expected.main_image_count,
                "K3": "副图数",
                "L3": expected.secondary_image_count,
                "A4": "店铺来源",
                "B4": STORE_URL,
                "E4": "排序",
                "F4": "Top Sales",
            },
            SHEET_NAMES[1]: {
                "A2": "UGREEN Shopee SKU 明细",
                "A3": "记录数",
                "B3": expected.sku_count,
                "E3": "商品数",
                "F3": expected.product_count,
                "A4": "说明",
                "B4": "本表列出详情页公开 SKU；商品当前价格见“商品汇总”。",
            },
            SHEET_NAMES[2]: {
                "A2": "UGREEN Shopee 商品图片明细",
                "A3": "图片记录数",
                "B3": expected.main_image_count + expected.secondary_image_count,
                "E3": "主图数",
                "F3": expected.main_image_count,
                "G3": "副图数",
                "H3": expected.secondary_image_count,
                "A4": "说明",
                "B4": "每个商品的主图列在前，副图保持详情页图库顺序。",
            },
            SHEET_NAMES[3]: {
                "A2": "UGREEN Shopee 全店抓取核验",
                "A3": "核验状态",
                "B3": expected.audit_status,
                "A4": "范围",
                "B4": "UGREEN Top Sales 全部分页与全部详情页",
            },
        }
        if known_details:
            expected_metadata[SHEET_NAMES[0]].update({"A2": KNOWN_SUMMARY_TITLE, "F4": KNOWN_ORDER})
            expected_metadata[SHEET_NAMES[1]]["B4"] = KNOWN_SKU_NOTE
            expected_metadata[SHEET_NAMES[3]].update({"A2": KNOWN_AUDIT_TITLE, "B4": KNOWN_SCOPE})
        for sheet_name, cells in expected_metadata.items():
            ws = wb[sheet_name]
            for coordinate, expected_value in cells.items():
                actual_value = ws[coordinate].value
                if actual_value != expected_value:
                    raise ValueError(
                        f"元数据 {sheet_name}!{coordinate} 不匹配："
                        f"应为 {expected_value!r}，实际为 {actual_value!r}"
                    )
        captured_at = summary["B3"].value
        if not isinstance(captured_at, str) or re.fullmatch(
            r"\d{4}/\d{2}/\d{2} \d{2}:\d{2}:\d{2}",
            captured_at,
        ) is None:
            raise ValueError("商品汇总!B3 的采集时间格式必须为 YYYY/MM/DD HH:MM:SS 文本")

        expected_rows = {
            SHEET_NAMES[0]: HEADER_ROW + expected.product_count,
            SHEET_NAMES[1]: HEADER_ROW + expected.sku_count,
            SHEET_NAMES[2]: HEADER_ROW + expected.main_image_count + expected.secondary_image_count,
            SHEET_NAMES[3]: HEADER_ROW + expected.audit_row_count,
        }
        for ws in wb.worksheets:
            if ws.max_row != expected_rows[ws.title]:
                raise ValueError(
                    f"工作表 {ws.title!r} 的行数不匹配："
                    f"应为 {expected_rows[ws.title]}，实际为 {ws.max_row}"
                )
            _assert_reference_layout(ws)
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        raise ValueError(f"在 {ws.title}!{cell.coordinate} 中发现公式")

        _assert_ids_are_text(
            summary,
            range(DATA_ROW, HEADER_ROW + expected.product_count + 1),
            (4, 5),
        )
        _assert_ids_are_text(
            sku,
            range(DATA_ROW, HEADER_ROW + expected.sku_count + 1),
            (3, 4, 6),
        )
        _assert_ids_are_text(
            images,
            range(DATA_ROW, HEADER_ROW + expected.main_image_count + expected.secondary_image_count + 1),
            (3, 4),
        )

        _assert_links(summary, (12, 14), expected.summary_hyperlink_count)
        _assert_links(sku, (8, 9), expected.sku_hyperlink_count)
        _assert_links(images, (8, 9), expected.image_hyperlink_count)

        audit_status = str(audit.cell(row=3, column=2).value or "")
        if audit_status != expected.audit_status:
            raise ValueError(
                f"核验状态不匹配：应为 {expected.audit_status!r}，实际为 {audit_status!r}"
            )
        return WorkbookStats(
            product_count=expected.product_count,
            sku_count=expected.sku_count,
            main_image_count=expected.main_image_count,
            secondary_image_count=expected.secondary_image_count,
            audit_status=audit_status,
            collection_mode=expected.collection_mode,
        )
    finally:
        wb.close()


def _required_int(
    value: Any,
    label: str,
    *,
    minimum: int = 0,
    maximum: int | None = None,
) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError(f"{label} 必须为整数，实际为 {value!r}")
    if not math.isfinite(float(value)) or int(value) != value:
        raise ValueError(f"{label} 必须为有限整数，实际为 {value!r}")
    result = int(value)
    if result < minimum or (maximum is not None and result > maximum):
        raise ValueError(f"{label} 超出允许范围：{result}")
    return result


def _required_text(value: Any, label: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label} 必须为非空文本")
    return value.strip()


def _required_http_url(value: Any, label: str) -> str:
    text = _required_text(value, label)
    if not _is_http_url(text):
        raise ValueError(f"{label} 不是 HTTP(S) URL：{text!r}")
    return text


def _pdp_identity(value: Any) -> tuple[str, str] | None:
    if not isinstance(value, str):
        return None
    try:
        parsed = urlsplit(value)
        hostname = (parsed.hostname or "").lower()
        if parsed.scheme.lower() != "https" or not (
            hostname == "shopee.ph" or hostname.endswith(".shopee.ph")
        ):
            return None
        path = unquote(parsed.path)
    except (TypeError, ValueError):
        return None
    for pattern in _PDP_IDENTITY_PATTERNS:
        if match := pattern.search(path):
            return str(int(match.group(1))), str(int(match.group(2)))
    return None


def _validate_store_url(value: Any) -> None:
    text = _required_text(value, "抓取核验/店铺 URL")
    try:
        parsed = urlsplit(text)
        query = parse_qs(parsed.query, keep_blank_values=True)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"抓取核验中的店铺 URL 无效：{text!r}") from exc
    if not (
        parsed.scheme.lower() == "https"
        and (parsed.hostname or "").lower() == "shopee.ph"
        and parsed.path.rstrip("/") == "/ugreen.ph"
        and query.get("page") == ["0"]
        and query.get("shop") == ["64922227"]
        and query.get("sortBy") == ["sales"]
        and query.get("tab") == ["0"]
    ):
        raise ValueError(f"抓取核验中的店铺 URL 与预期不符：{text!r}")


def validate_excel(path: str | os.PathLike[str]) -> WorkbookStats:
    """Reopen and independently validate the canonical workbook and its joins."""

    workbook_path = Path(path)
    _verify_zip(workbook_path)
    wb = load_workbook(workbook_path, data_only=False, read_only=False, keep_links=True)
    try:
        if tuple(wb.sheetnames) != SHEET_NAMES:
            raise ValueError(f"工作表顺序异常：{wb.sheetnames!r}")
        summary = wb[SHEET_NAMES[0]]
        sku = wb[SHEET_NAMES[1]]
        images = wb[SHEET_NAMES[2]]
        audit = wb[SHEET_NAMES[3]]
        audit_values: dict[str, Any] = {}
        for row in range(DATA_ROW, audit.max_row + 1):
            label = _required_text(audit.cell(row, 1).value, f"抓取核验!A{row}")
            if label in audit_values:
                raise ValueError(f"抓取核验项重复：{label!r}")
            audit_values[label] = audit.cell(row, 2).value
        collection_mode = audit_values.get("采集模式", FULL_TOPSALES_MODE)
        if collection_mode not in {FULL_TOPSALES_MODE, KNOWN_DETAILS_MODE}:
            raise ValueError("抓取核验中的采集模式无效")
        known_details = collection_mode == KNOWN_DETAILS_MODE
        _assert_headers(summary, KNOWN_SUMMARY_HEADERS if known_details else SUMMARY_HEADERS)
        _assert_headers(sku, KNOWN_SKU_HEADERS if known_details else SKU_HEADERS)
        _assert_headers(images, KNOWN_IMAGE_HEADERS if known_details else IMAGE_HEADERS)
        _assert_headers(audit, AUDIT_HEADERS)
        for ws in wb.worksheets:
            _assert_reference_layout(ws)
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        raise ValueError(f"在 {ws.title}!{cell.coordinate} 中发现公式")

        product_count = summary.max_row - HEADER_ROW
        sku_count = sku.max_row - HEADER_ROW
        if product_count < 1 or sku_count < 1:
            raise ValueError("工作簿中没有商品或 SKU 记录")
        fixed_metadata = {
            (SHEET_NAMES[0], "A2"): "UGREEN Shopee Top Sales 全店商品",
            (SHEET_NAMES[0], "A3"): "采集时间",
            (SHEET_NAMES[0], "E3"): "商品数",
            (SHEET_NAMES[0], "F3"): product_count,
            (SHEET_NAMES[0], "G3"): "SKU数",
            (SHEET_NAMES[0], "H3"): sku_count,
            (SHEET_NAMES[0], "A4"): "店铺来源",
            (SHEET_NAMES[0], "B4"): STORE_URL,
            (SHEET_NAMES[0], "E4"): "排序",
            (SHEET_NAMES[0], "F4"): "Top Sales",
            (SHEET_NAMES[1], "A2"): "UGREEN Shopee SKU 明细",
            (SHEET_NAMES[1], "A3"): "记录数",
            (SHEET_NAMES[1], "B3"): sku_count,
            (SHEET_NAMES[1], "E3"): "商品数",
            (SHEET_NAMES[1], "F3"): product_count,
            (SHEET_NAMES[2], "A2"): "UGREEN Shopee 商品图片明细",
            (SHEET_NAMES[2], "A3"): "图片记录数",
            (SHEET_NAMES[2], "B3"): images.max_row - HEADER_ROW,
            (SHEET_NAMES[3], "A2"): "UGREEN Shopee 全店抓取核验",
            (SHEET_NAMES[3], "A3"): "核验状态",
            (SHEET_NAMES[3], "B3"): "通过",
            (SHEET_NAMES[3], "A4"): "范围",
            (SHEET_NAMES[3], "B4"): "UGREEN Top Sales 全部分页与全部详情页",
        }
        if known_details:
            fixed_metadata.update({
                (SHEET_NAMES[0], "A2"): KNOWN_SUMMARY_TITLE,
                (SHEET_NAMES[0], "F4"): KNOWN_ORDER,
                (SHEET_NAMES[1], "B4"): KNOWN_SKU_NOTE,
                (SHEET_NAMES[3], "A2"): KNOWN_AUDIT_TITLE,
                (SHEET_NAMES[3], "B4"): KNOWN_SCOPE,
            })
        for (sheet_name, coordinate), expected_value in fixed_metadata.items():
            actual_value = wb[sheet_name][coordinate].value
            if actual_value != expected_value:
                raise ValueError(
                    f"元数据 {sheet_name}!{coordinate} 不匹配："
                    f"应为 {expected_value!r}，实际为 {actual_value!r}"
                )
        captured_at = summary["B3"].value
        if not isinstance(captured_at, str) or re.fullmatch(
            r"\d{4}/\d{2}/\d{2} \d{2}:\d{2}:\d{2}",
            captured_at,
        ) is None:
            raise ValueError("商品汇总!B3 的采集时间格式必须为 YYYY/MM/DD HH:MM:SS 文本")
        _validate_store_url(summary["B4"].value)
        _assert_ids_are_text(
            summary,
            range(DATA_ROW, summary.max_row + 1),
            (4, 5),
        )
        _assert_ids_are_text(
            sku,
            range(DATA_ROW, sku.max_row + 1),
            (3, 4, 6),
        )
        _assert_ids_are_text(
            images,
            range(DATA_ROW, images.max_row + 1),
            (3, 4),
        )

        products: dict[tuple[str, str], dict[str, Any]] = {}
        pages: set[int] = set()
        missing_monthly = 0
        for row in range(DATA_ROW, summary.max_row + 1):
            rank = _required_int(summary.cell(row, 1).value, f"商品汇总!A{row}", minimum=1)
            if rank != row - HEADER_ROW:
                raise ValueError(f"商品汇总!A{row} 中的商品排名不连续")
            if known_details:
                page = None
                for column in (2, 3, 7, 8, 9):
                    if summary.cell(row, column).value is not None:
                        raise ValueError(f"商品汇总第 {row} 行含有本次未采集的列表字段")
            else:
                page = _required_int(summary.cell(row, 2).value, f"商品汇总!B{row}", minimum=1)
                _required_int(summary.cell(row, 3).value, f"商品汇总!C{row}", minimum=1, maximum=30)
                pages.add(page)
            shop_id = _required_text(summary.cell(row, 4).value, f"商品汇总!D{row}")
            item_id = _required_text(summary.cell(row, 5).value, f"商品汇总!E{row}")
            if shop_id != "64922227" or not item_id.isdigit():
                raise ValueError(f"商品汇总第 {row} 行的商品标识无效")
            key = (shop_id, item_id)
            if key in products:
                raise ValueError(f"商品汇总第 {row} 行存在重复商品标识：{key!r}")
            title = _required_text(summary.cell(row, 6).value, f"商品汇总!F{row}")
            price = summary.cell(row, 7).value
            if not known_details and (
                isinstance(price, bool)
                or not isinstance(price, (int, float, Decimal))
                or not math.isfinite(float(price))
                or price < 0
            ):
                raise ValueError(f"商品汇总!G{row} 中的商品价格无效：{price!r}")
            monthly_status = summary.cell(row, 10).value
            if monthly_status not in ({"未采集"} if known_details else {"已展示", "未展示"}):
                raise ValueError(f"商品汇总!J{row} 中的月销状态无效")
            if monthly_status == "未展示":
                missing_monthly += 1
                if summary.cell(row, 8).value not in (None, "") or summary.cell(row, 9).value not in (None, ""):
                    raise ValueError(f"商品汇总第 {row} 行未展示的月销数据必须留空")
            elif not known_details:
                _required_text(summary.cell(row, 8).value, f"商品汇总!H{row}")
                _required_int(summary.cell(row, 9).value, f"商品汇总!I{row}")
            expected_skus = _required_int(
                summary.cell(row, 11).value,
                f"商品汇总!K{row}",
                minimum=1,
            )
            main_url = _required_http_url(summary.cell(row, 12).value, f"商品汇总!L{row}")
            expected_secondary = _required_int(
                summary.cell(row, 13).value,
                f"商品汇总!M{row}",
            )
            product_url = _required_http_url(summary.cell(row, 14).value, f"商品汇总!N{row}")
            if _pdp_identity(product_url) != key:
                raise ValueError(f"商品汇总!N{row} 中的 PDP URL 商品标识不匹配")
            products[key] = {
                "rank": rank,
                "page": page,
                "title": title,
                "product_url": product_url,
                "main_url": main_url,
                "expected_skus": expected_skus,
                "expected_secondary": expected_secondary,
            }
        if not known_details and pages != set(range(1, max(pages) + 1)):
            raise ValueError(f"商品汇总包含不连续的列表页码：{sorted(pages)!r}")

        sku_per_product = {key: 0 for key in products}
        model_keys: set[tuple[str, str, str]] = set()
        for row in range(DATA_ROW, sku.max_row + 1):
            rank = _required_int(sku.cell(row, 1).value, f"SKU明细!A{row}", minimum=1)
            page = sku.cell(row, 2).value if known_details else _required_int(sku.cell(row, 2).value, f"SKU明细!B{row}", minimum=1)
            shop_id = _required_text(sku.cell(row, 3).value, f"SKU明细!C{row}")
            item_id = _required_text(sku.cell(row, 4).value, f"SKU明细!D{row}")
            key = (shop_id, item_id)
            product = products.get(key)
            if product is None:
                raise ValueError(f"SKU明细第 {row} 行引用了未知商品")
            if rank != product["rank"] or page != product["page"]:
                raise ValueError(f"SKU明细第 {row} 行的排名或列表页不匹配")
            if _required_text(sku.cell(row, 5).value, f"SKU明细!E{row}") != product["title"]:
                raise ValueError(f"SKU明细第 {row} 行的商品标题不匹配")
            model_id = _required_text(sku.cell(row, 6).value, f"SKU明细!F{row}")
            model_key = (shop_id, item_id, model_id)
            if not model_id.isdigit() or model_key in model_keys:
                raise ValueError(f"SKU明细第 {row} 行的 model ID 无效或重复")
            model_keys.add(model_key)
            _required_text(sku.cell(row, 7).value, f"SKU明细!G{row}")
            image_url = sku.cell(row, 8).value
            if image_url not in (None, ""):
                _required_http_url(image_url, f"SKU明细!H{row}")
            if _required_http_url(sku.cell(row, 9).value, f"SKU明细!I{row}") != product["product_url"]:
                raise ValueError(f"SKU明细第 {row} 行的 PDP URL 不匹配")
            sku_per_product[key] += 1
        for key, product in products.items():
            if sku_per_product[key] != product["expected_skus"]:
                raise ValueError(f"商品 {key!r} 的 SKU 数量不匹配")

        main_urls: dict[tuple[str, str], list[str]] = {key: [] for key in products}
        secondary_indices: dict[tuple[str, str], list[int]] = {key: [] for key in products}
        secondary_count = 0
        for row in range(DATA_ROW, images.max_row + 1):
            rank = _required_int(images.cell(row, 1).value, f"图片明细!A{row}", minimum=1)
            page = images.cell(row, 2).value if known_details else _required_int(images.cell(row, 2).value, f"图片明细!B{row}", minimum=1)
            shop_id = _required_text(images.cell(row, 3).value, f"图片明细!C{row}")
            item_id = _required_text(images.cell(row, 4).value, f"图片明细!D{row}")
            key = (shop_id, item_id)
            product = products.get(key)
            if product is None:
                raise ValueError(f"图片明细第 {row} 行引用了未知商品")
            if rank != product["rank"] or page != product["page"]:
                raise ValueError(f"图片明细第 {row} 行的排名或列表页不匹配")
            if _required_text(images.cell(row, 5).value, f"图片明细!E{row}") != product["title"]:
                raise ValueError(f"图片明细第 {row} 行的商品标题不匹配")
            image_type = images.cell(row, 6).value
            image_index = _required_int(images.cell(row, 7).value, f"图片明细!G{row}", minimum=1)
            image_url = _required_http_url(images.cell(row, 8).value, f"图片明细!H{row}")
            if _required_http_url(images.cell(row, 9).value, f"图片明细!I{row}") != product["product_url"]:
                raise ValueError(f"图片明细第 {row} 行的 PDP URL 不匹配")
            if image_type == "主图":
                if image_index != 1:
                    raise ValueError(f"图片明细第 {row} 行的主图序号必须为 1")
                main_urls[key].append(image_url)
            elif image_type == "副图":
                secondary_count += 1
                secondary_indices[key].append(image_index)
            else:
                raise ValueError(f"图片明细!F{row} 中存在未知图片类型：{image_type!r}")
        for key, product in products.items():
            if main_urls[key] != [product["main_url"]]:
                raise ValueError(f"商品 {key!r} 的主图不匹配")
            expected_indices = list(range(1, product["expected_secondary"] + 1))
            if sorted(secondary_indices[key]) != expected_indices:
                raise ValueError(f"商品 {key!r} 的副图序列不匹配")

        count_metadata = {
            (SHEET_NAMES[0], "I3"): "主图数",
            (SHEET_NAMES[0], "J3"): product_count,
            (SHEET_NAMES[0], "K3"): "副图数",
            (SHEET_NAMES[0], "L3"): secondary_count,
            (SHEET_NAMES[2], "E3"): "主图数",
            (SHEET_NAMES[2], "F3"): product_count,
            (SHEET_NAMES[2], "G3"): "副图数",
            (SHEET_NAMES[2], "H3"): secondary_count,
        }
        for (sheet_name, coordinate), expected_value in count_metadata.items():
            actual_value = wb[sheet_name][coordinate].value
            if actual_value != expected_value:
                raise ValueError(
                    f"元数据 {sheet_name}!{coordinate} 不匹配："
                    f"应为 {expected_value!r}，实际为 {actual_value!r}"
                )

        expected_audit_labels = (
            "店铺 URL",
            "排序方式",
            "列表总页数",
            "商品卡出现次数",
            "跨页重复次数",
            "唯一商品数",
            "成功进入详情页",
            "详情页失败数",
            "详情解析成功数",
            "详情解析错误数",
            "未展示月销商品数",
            "SKU 总数",
            "有 SKU 的商品数",
            "主图数",
            "副图数",
            "无效图片 URL 数",
            "完整性结论",
        )
        if known_details:
            expected_audit_labels = (
                "店铺 URL", "采集模式", "参考工作簿", "参考日期", "参考采集时间",
                "参考文件 SHA-256", "参考身份序列 SHA-256", "参考唯一商品数",
                "用户确认无新增商品", "本次重抓列表", "排序方式", "列表总页数",
                "商品卡出现次数", "跨页重复次数", "唯一商品数", "成功进入详情页",
                "详情页失败数", "详情解析成功数", "详情解析错误数", "未采集月销商品数",
                "SKU 总数", "有 SKU 的商品数", "主图数", "副图数", "无效图片 URL 数", "完整性结论",
            )
        if tuple(audit_values) != expected_audit_labels:
            raise ValueError(f"抓取核验项目或顺序异常：{tuple(audit_values)!r}")
        audit_status = str(audit.cell(row=3, column=2).value or "")
        if audit_status != "通过" or audit_values.get("完整性结论") != "通过":
            raise ValueError("工作簿核验状态不是“通过”")
        if audit_values.get("排序方式") != (KNOWN_ORDER if known_details else "Top Sales"):
            raise ValueError("工作簿核验排序方式与采集模式不匹配")
        _validate_store_url(audit_values.get("店铺 URL"))
        expected_audit_counts = {
            "列表总页数": 0 if known_details else max(pages),
            "唯一商品数": product_count,
            "成功进入详情页": product_count,
            "详情页失败数": 0,
            "详情解析成功数": product_count,
            "详情解析错误数": 0,
            "未展示月销商品数": missing_monthly,
            "SKU 总数": sku_count,
            "有 SKU 的商品数": product_count,
            "主图数": product_count,
            "副图数": secondary_count,
            "无效图片 URL 数": 0,
        }
        if known_details:
            del expected_audit_counts["未展示月销商品数"]
            expected_audit_counts.update({
                "未采集月销商品数": product_count,
                "参考唯一商品数": product_count,
                "商品卡出现次数": 0,
                "跨页重复次数": 0,
            })
            if audit_values.get("用户确认无新增商品") != "是" or audit_values.get("本次重抓列表") != "否":
                raise ValueError("已知商品详情刷新的范围确认或列表采集声明无效")
            _validate_known_scope({
                "reference_workbook": audit_values.get("参考工作簿"),
                "reference_date": audit_values.get("参考日期"),
                "reference_captured_at": audit_values.get("参考采集时间"),
                "reference_sha256": audit_values.get("参考文件 SHA-256"),
                "reference_identity_sha256": audit_values.get("参考身份序列 SHA-256"),
                "reference_unique_count": audit_values.get("参考唯一商品数"),
                "user_confirmed_no_new_products": True,
            }, list(products), reference_name_only=True)
            if audit_values["参考采集时间"] > captured_at:
                raise ValueError("参考采集时间晚于本次详情采集时间")
            if audit.cell(audit.max_row, 3).value != KNOWN_CONCLUSION:
                raise ValueError("已知商品详情刷新完整性结论未明确限定参考清单范围")
        for label, expected in expected_audit_counts.items():
            actual = _required_int(audit_values.get(label), f"抓取核验/{label}")
            if actual != expected:
                raise ValueError(
                    f"抓取核验项 {label!r} 的数量不匹配：应为 {expected}，实际为 {actual}"
                )
        occurrences = _required_int(audit_values.get("商品卡出现次数"), "抓取核验/商品卡出现次数", minimum=0 if known_details else 1)
        duplicates = _required_int(audit_values.get("跨页重复次数"), "抓取核验/跨页重复次数")
        if not known_details and occurrences - duplicates != product_count:
            raise ValueError("抓取核验中的商品卡出现次数与去重次数无法核对一致")
        return WorkbookStats(
            product_count=product_count,
            sku_count=sku_count,
            main_image_count=product_count,
            secondary_image_count=secondary_count,
            audit_status=audit_status,
            collection_mode=collection_mode,
        )
    finally:
        wb.close()


def write_excel(result: Any, output_path: str | os.PathLike[str]) -> Path:
    """Write, reopen, verify and atomically publish one ``.xlsx`` workbook.

    No JSON, HTML, image preview or validation sidecar is created.  If either
    saving or validation fails, a pre-existing destination workbook is left
    untouched and the temporary file is removed.
    """

    destination = Path(output_path).expanduser()
    if destination.suffix.lower() != ".xlsx":
        raise ValueError(f"Excel 输出文件必须使用 .xlsx 扩展名：{destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)

    view = _normalise_export(result)
    workbook, expected = _build_workbook(view)
    if expected.audit_status != "通过":
        workbook.close()
        raise ValueError("抓取完整性校验失败，拒绝替换 Excel 文件")
    try:
        file_descriptor, temporary_name = tempfile.mkstemp(
            dir=destination.parent,
            prefix=f".{destination.stem}.",
            suffix=".xlsx",
        )
    except BaseException:
        workbook.close()
        raise
    temporary_path = Path(temporary_name)
    descriptor_open = True
    try:
        os.close(file_descriptor)
        descriptor_open = False
        try:
            workbook.save(temporary_path)
        finally:
            workbook.close()
        with temporary_path.open("rb") as handle:
            os.fsync(handle.fileno())
        _validate_with_expected(temporary_path, expected)
        if view.collection_mode == KNOWN_DETAILS_MODE:
            validate_excel(temporary_path)
        os.replace(temporary_path, destination)
        try:
            directory_descriptor = os.open(destination.parent, os.O_RDONLY)
            try:
                os.fsync(directory_descriptor)
            finally:
                os.close(directory_descriptor)
        except OSError:
            # Directory fsync is not available on every platform/filesystem.
            pass
        return destination
    except BaseException:
        if descriptor_open:
            try:
                os.close(file_descriptor)
            except OSError:
                pass
        temporary_path.unlink(missing_ok=True)
        raise
