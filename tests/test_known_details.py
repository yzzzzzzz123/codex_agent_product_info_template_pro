"""已知清单详情刷新离线契约；不启动浏览器、不联网、不读真实 profile。"""

from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
from contextlib import ExitStack
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/shopee-ugreen-topsales/scripts"))
import scraper  # noqa: E402
from test_scraper import FIXTURE_VERSION, fingerprint

US_REGION = scraper.BrowserRegion("US", "en-US", ("en-US", "en"), "America/New_York")


def make_reference(path: Path, ids: tuple[str, ...] = ("100", "200", "300")) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品汇总"
    sheet["A3"] = "采集时间"
    sheet["B3"] = "2026/09/17 15:00:00"
    headers = ("店铺ID", "商品ID", "商品链接", "列表页", "价格PHP", "商品标题", "月销")
    for index, value in enumerate(headers, 1):
        sheet.cell(6, index, value)
    for item_id in ids:
        sheet.append([
            scraper.SHOP_ID, item_id,
            f"https://shopee.ph/product/{scraper.SHOP_ID}/{item_id}",
            "历史页不应进入新结果", 987654, "历史标题不得复用", 876543,
        ])
    workbook.save(path)
    workbook.close()


def known_card(item_id: str = "100") -> scraper.ProductCard:
    return scraper.ProductCard(
        None, None, scraper.SHOP_ID, item_id, "",
        f"https://shopee.ph/product/{scraper.SHOP_ID}/{item_id}",
        None, None, None, None, "not_collected",
    )


def payload(item_id: str = "100") -> dict:
    return {"bff": {"item": {
        "shop_id": scraper.SHOP_ID, "item_id": item_id,
        "title": f"实时商品 {item_id}",
        "price": 9999900000, "sold": 888888, "historical_sold": 777777,
        "models": [{"model_id": "12345"}],
        "images": ["current-main", "current-secondary"],
    }}}


class ReferenceTests(unittest.TestCase):
    def test_entire_reference_is_deduplicated_in_order_without_old_fields(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.xlsx"
            make_reference(path, ("100", "200", "100", "300"))
            before = path.read_bytes()
            cards, metadata = scraper._read_known_product_reference(path)
            self.assertEqual([card.item_id for card in cards], ["100", "200", "300"])
            for card in cards:
                self.assertIsNone(card.source_page)
                self.assertIsNone(card.source_position)
                self.assertIsNone(card.price_php)
                self.assertIsNone(card.monthly_sales_text)
                self.assertIsNone(card.monthly_sales_display)
                self.assertIsNone(card.monthly_sales_count_lower_bound)
                self.assertEqual(card.title, "")
                self.assertEqual(card.monthly_sales_observation, "not_collected")
            self.assertEqual(metadata["reference_sha256"], hashlib.sha256(before).hexdigest())
            self.assertEqual(metadata["reference_workbook"], str(path.resolve()))
            self.assertEqual(metadata["reference_date"], "2026-09-17")
            self.assertEqual(metadata["reference_captured_at"], "2026/09/17 15:00:00")
            self.assertEqual(metadata["reference_unique_count"], 3)
            identities = "\n".join(f"{scraper.SHOP_ID}:{value}" for value in ("100", "200", "300"))
            self.assertEqual(metadata["reference_identity_sha256"], hashlib.sha256(identities.encode()).hexdigest())
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(list(Path(directory).iterdir()), [path])

    def test_missing_or_invalid_date_is_not_replaced_with_mtime(self) -> None:
        from openpyxl import load_workbook
        for stamp in (None, "yesterday", "2026/02/30 10:00:00", "2026/9/17 15:00:00"):
            with self.subTest(stamp=stamp), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "fixture.xlsx"
                make_reference(path)
                workbook = load_workbook(path)
                workbook["商品汇总"]["B3"] = stamp
                workbook.save(path)
                workbook.close()
                with self.assertRaisesRegex(ValueError, "采集时间"):
                    scraper._read_known_product_reference(path)

    def test_invalid_foreign_mismatched_or_numeric_identity_rejected(self) -> None:
        from openpyxl import load_workbook
        cases = (
            ("A7", "111"), ("B7", 100), ("B7", "00100"),
            ("C7", "https://example.com/product/64922227/100"),
            ("C7", "http://shopee.ph/product/64922227/100"),
            ("C7", "https://shopee.ph/product/64922227/999"),
            ("C7", "https://user@shopee.ph/product/64922227/100"),
            ("C7", "https://shopee.ph:999/product/64922227/100"),
            ("B7", '=CONCAT("1","00")'),
        )
        for cell, value in cases:
            with self.subTest(cell=cell, value=value), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "fixture.xlsx"
                make_reference(path)
                workbook = load_workbook(path)
                workbook["商品汇总"][cell] = value
                workbook.save(path)
                workbook.close()
                with self.assertRaises(ValueError):
                    scraper._read_known_product_reference(path)

    def test_empty_reference_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.xlsx"
            make_reference(path, ())
            with self.assertRaisesRegex(ValueError, "没有已知商品"):
                scraper._read_known_product_reference(path)


class KnownParsingTests(unittest.TestCase):
    def test_title_skus_and_gallery_are_current_but_list_fields_stay_uncollected(self) -> None:
        card = known_card()
        card.title = "故意污染的旧标题"
        result = scraper._parse_detail(card, payload(), 12)
        self.assertEqual(result.title, "实时商品 100")
        self.assertEqual([sku.model_id for sku in result.skus], ["12345"])
        self.assertTrue(result.main_image_url.endswith("current-main"))
        self.assertTrue(result.secondary_image_urls[0].endswith("current-secondary"))
        for field in ("global_rank", "source_page", "source_position", "price_php",
                      "monthly_sales_text", "monthly_sales_count_lower_bound"):
            self.assertIsNone(getattr(result, field), field)
        self.assertEqual(result.monthly_sales_observation, "not_collected")

    def test_missing_nontext_or_blank_current_title_rejected(self) -> None:
        for title in (None, "", "   ", 100):
            data = payload()
            data["bff"]["item"]["title"] = title
            with self.subTest(title=title), self.assertRaisesRegex(scraper.ScrapeError, "标题"):
                scraper._parse_detail(known_card(), data, 1)


class KnownRefreshFlowTests(unittest.IsolatedAsyncioTestCase):
    async def test_explicit_user_confirmation_required_before_any_reference_or_browser_read(self) -> None:
        with patch.object(scraper, "_read_known_product_reference") as read, \
                patch.object(scraper, "_temporary_chrome_profile_base") as profile:
            for confirmation in (False, None, 1, "true"):
                with self.subTest(confirmation=confirmation), self.assertRaisesRegex(ValueError, "确认"):
                    await scraper.scrape_known_details(None, Path("unused"),
                                                       no_new_products_confirmed=confirmation)
            read.assert_not_called()
            profile.assert_not_called()

    async def test_multichannel_rejected(self) -> None:
        config = scraper.BrowserConfig(browser_profile="windows-intel", chrome_version=FIXTURE_VERSION, region=US_REGION, detail_shards=2, chrome_executable=Path(__file__))
        with self.assertRaisesRegex(ValueError, "单路"):
            await scraper.scrape_known_details(config, Path("unused"), no_new_products_confirmed=True)

    async def run_refresh(self, mutate=None, failure=None):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.xlsx"
            make_reference(path)
            before = path.read_bytes()
            profile = MagicMock()
            base = Path("/unused-offline-profile")
            profile.__enter__.return_value = base
            profile.__exit__.return_value = False
            context = object()
            config = scraper.BrowserConfig(browser_profile="windows-intel", chrome_version=FIXTURE_VERSION, region=US_REGION, chrome_executable=Path(__file__), historical_list_preflight=True)

            async def collect(contexts, cards, passed_config, progress):
                self.assertEqual(contexts, [context])
                self.assertIs(passed_config, config)
                self.assertEqual([card.item_id for card in cards], ["100", "200", "300"])
                if failure:
                    raise failure
                products = [scraper._parse_detail(card, payload(card.item_id), rank)
                            for rank, card in enumerate(cards, 1)]
                return mutate(products) if mutate else products

            with ExitStack() as stack:
                stack.enter_context(patch.object(scraper, "_temporary_chrome_profile_base", return_value=profile))
                launch = stack.enter_context(patch.object(scraper, "_launch_context", new=AsyncMock(return_value=context)))
                close = stack.enter_context(patch.object(scraper, "_close_context", new_callable=AsyncMock))
                stack.enter_context(patch.object(scraper, "_collect_details", side_effect=collect))
                for name in ("_collect_listing", "_historical_list_preflight", "_capture_list_page_with_profile",
                             "_capture_node_list_snapshot", "_clone_detail_profiles"):
                    stack.enter_context(patch.object(scraper, name, side_effect=AssertionError("不得使用列表链路")))
                stack.enter_context(patch("excel.write_excel", side_effect=AssertionError("抓取器不导出")))
                try:
                    result = await scraper.scrape_known_details(config, path, no_new_products_confirmed=True)
                finally:
                    launch.assert_awaited_once_with(config, base)
                    close.assert_awaited_once_with(context)
                    profile.__exit__.assert_called_once()
                    self.assertEqual(path.read_bytes(), before)
                    self.assertEqual(list(Path(directory).iterdir()), [path])
                return result

    async def test_all_known_products_refreshed_without_list_or_old_data(self) -> None:
        result = await self.run_refresh()
        self.assertEqual(result.collection_mode, "known_product_details")
        self.assertEqual([product.item_id for product in result.products], ["100", "200", "300"])
        self.assertEqual(result.audit.detail_success_count, 3)
        self.assertEqual(result.audit.detail_failure_count, 0)
        self.assertEqual(result.audit.listed_product_count, 3)
        self.assertEqual(result.audit.list_page_count, 0)
        self.assertEqual(result.audit.list_input_product_occurrences, 0)
        self.assertEqual(result.audit.list_duplicate_occurrence_count, 0)
        self.assertEqual(result.audit.products_without_monthly_sales_display, 0)
        self.assertTrue(result.scope_metadata["user_confirmed_no_new_products"])
        self.assertEqual(result.scope_metadata["reference_unique_count"], 3)

    async def test_missing_extra_duplicate_or_reordered_results_rejected(self) -> None:
        mutations = (
            lambda rows: rows[:-1], lambda rows: rows + rows[:1],
            lambda rows: rows[:1] + rows[:1] + rows[2:], lambda rows: list(reversed(rows)),
        )
        for mutation in mutations:
            with self.subTest(mutation=mutation), self.assertRaisesRegex(scraper.ScrapeError, "完整性"):
                await self.run_refresh(mutate=mutation)

    async def test_uncollected_fields_cannot_be_backfilled(self) -> None:
        for field, value in (("global_rank", 1), ("source_page", 0), ("source_position", 1),
                             ("price_php", 1), ("monthly_sales_text", "1 Sold/Month"),
                             ("monthly_sales_count_lower_bound", 1),
                             ("monthly_sales_observation", "not_displayed"), ("title", "")):
            with self.subTest(field=field), self.assertRaisesRegex(scraper.ScrapeError, "混入"):
                await self.run_refresh(mutate=lambda rows: [replace(rows[0], **{field: value})] + rows[1:])

    async def test_detail_failure_returns_no_result_and_closes_context(self) -> None:
        with self.assertRaisesRegex(scraper.AccessChallengeError, "拒绝"):
            await self.run_refresh(failure=scraper.AccessChallengeError("详情拒绝"))


class DetailsOnlyVerificationTests(unittest.IsolatedAsyncioTestCase):
    async def test_first_middle_last_sample_skips_every_list_path(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.xlsx"
            make_reference(path, ("100", "200", "300", "400", "500"))
            page = SimpleNamespace(close=AsyncMock())
            context = SimpleNamespace(new_page=AsyncMock(return_value=page))
            profile = MagicMock()
            profile.__enter__.return_value = Path("/unused-offline-profile")
            profile.__exit__.return_value = False

            async def visit(_page, card, rank, _config):
                return scraper._parse_detail(card, payload(card.item_id), rank)

            with ExitStack() as stack:
                stack.enter_context(patch.object(scraper, "_temporary_chrome_profile_base", return_value=profile))
                stack.enter_context(patch.object(scraper, "_launch_context", new=AsyncMock(return_value=context)))
                close = stack.enter_context(patch.object(scraper, "_close_context", new_callable=AsyncMock))
                visited = stack.enter_context(patch.object(scraper, "_visit_detail", side_effect=visit))
                park = stack.enter_context(patch.object(scraper, "_park_and_wait", new_callable=AsyncMock))
                sleep = stack.enter_context(patch.object(scraper.asyncio, "sleep", new_callable=AsyncMock))
                for name in ("_select_verification_targets", "_historical_list_preflight",
                             "_capture_list_page_with_profile", "_capture_node_list_snapshot"):
                    stack.enter_context(patch.object(scraper, name, side_effect=AssertionError("不得使用列表链路")))
                config = scraper.BrowserConfig(browser_profile="windows-intel", chrome_version=FIXTURE_VERSION, region=US_REGION, chrome_executable=Path(__file__), historical_list_preflight=True)
                report = await scraper.verify_access(config, path, details_only=True)
                self.assertEqual([call.args[1].item_id for call in visited.await_args_list], ["100", "300", "500"])
                self.assertEqual([call.args[1] for call in park.await_args_list], [10000, 10000, 0])
                sleep.assert_not_awaited()
                close.assert_awaited_once_with(context)
                page.close.assert_awaited_once()
            self.assertEqual(report["mode"], "verify_access")
            self.assertTrue(report["details_only"])
            self.assertTrue(report["list_skipped"])
            self.assertIsNone(report["list_pass"])
            self.assertEqual(report["list_pages_checked"], 0)
            self.assertFalse(report["historical_list_preflight"])
            self.assertIsNone(report["preflight_pass"])
            self.assertEqual(report["scope_metadata"]["reference_unique_count"], 5)
            self.assertEqual(report["detail_success_count"], 3)
            self.assertEqual(report["detail_target_count"], 3)
            self.assertFalse(report["full_crawl_completed"])
            self.assertFalse(report["workbook_written"])


if __name__ == "__main__":
    unittest.main()
