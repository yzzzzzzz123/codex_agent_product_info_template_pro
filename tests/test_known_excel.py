"""离线检查已知商品详情模式；所有虚构夹具仅位于自动清理的临时目录。"""

from __future__ import annotations

from copy import deepcopy
import hashlib
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "skills/shopee-ugreen-topsales/scripts"
sys.path.insert(0, str(SCRIPTS))
import excel  # noqa: E402


def fixture(*, known=True):
    products = []
    for index in (1, 2):
        products.append({
            "global_rank": None if known else index,
            "source_page": None if known else 0,
            "source_position": None if known else index,
            "shop_id": "64922227", "item_id": str(100 + index),
            "title": f"离线夹具商品 {index}",
            "product_url": f"https://shopee.ph/product/64922227/{100 + index}",
            "price_php": None if known else 10,
            "monthly_sales_text": None,
            "monthly_sales_count_lower_bound": None,
            "monthly_sales_observation": "not_collected" if known else "not_displayed",
            "skus": [{"model_id": str(1000 + index), "sku_props": "离线规格",
                      "sku_image_url": f"https://example.invalid/sku-{index}.png"}],
            "main_image_url": f"https://example.invalid/main-{index}.png",
            "secondary_image_urls": [f"https://example.invalid/secondary-{index}.png"],
        })
    result = {
        "captured_at": "2000/01/02 12:00:00", "products": products,
        "audit": {"list_page_count": 0 if known else 1,
                  "list_input_product_occurrences": 0 if known else 2,
                  "list_duplicate_occurrence_count": 0,
                  "listed_product_count": 2, "detail_success_count": 2,
                  "detail_failure_count": 0,
                  "products_without_monthly_sales_display": 0 if known else 2},
    }
    if known:
        result.update({
            "collection_mode": "known_product_details",
            "scope_metadata": {
                "reference_workbook": "/offline-fixture/reference.xlsx",
                "reference_sha256": "a" * 64,
                "reference_captured_at": "2000/01/01 12:00:00",
                "reference_date": "2000-01-01", "reference_unique_count": 2,
                "reference_identity_sha256": hashlib.sha256(b"64922227:101\n64922227:102").hexdigest(),
                "user_confirmed_no_new_products": True,
            },
        })
    return result


class KnownExcelTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="ugreen-excel-test-")
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        self.path = self.directory / "offline_fixture.xlsx"

    def publish(self, result=None):
        excel.write_excel(fixture() if result is None else result, self.path)

    def mutate(self, callback):
        wb = load_workbook(self.path)
        try:
            callback(wb)
            wb.save(self.path)
        finally:
            wb.close()

    @staticmethod
    def audit_row(wb, label):
        ws = wb["抓取核验"]
        return next(row for row in range(excel.DATA_ROW, ws.max_row + 1) if ws.cell(row, 1).value == label)

    def test_known_roundtrip_preserves_layout_and_marks_uncollected(self):
        self.publish()
        stats = excel.validate_excel(self.path)
        self.assertEqual((stats.product_count, stats.sku_count, stats.main_image_count,
                          stats.secondary_image_count, stats.audit_status, stats.collection_mode),
                         (2, 2, 2, 2, "通过", "known_product_details"))
        wb = load_workbook(self.path)
        try:
            self.assertEqual(tuple(wb.sheetnames), excel.SHEET_NAMES)
            summary = wb["商品汇总"]
            self.assertEqual(summary["A6"].value, "采集序号")
            self.assertEqual(summary["A7"].value, 1)
            for column in (2, 3, 7, 8, 9):
                self.assertIsNone(summary.cell(7, column).value)
            self.assertEqual(summary["J7"].value, "未采集")
            self.assertEqual(summary.freeze_panes, "F7")
            self.assertEqual(summary["D7"].number_format, "@")
            self.assertEqual(wb["SKU明细"].freeze_panes, "E7")
            self.assertEqual(wb["图片明细"].freeze_panes, "E7")
            self.assertEqual(wb["抓取核验"].freeze_panes, "A7")
            audit = wb["抓取核验"]
            row = self.audit_row(wb, "参考工作簿")
            self.assertEqual(audit.cell(row, 2).value, "reference.xlsx")
            self.assertEqual(audit.cell(audit.max_row, 3).value, excel.KNOWN_CONCLUSION)
        finally:
            wb.close()
        self.assertEqual(list(self.directory.iterdir()), [self.path])

    def test_original_full_mode_schema_stays_unchanged(self):
        self.publish(fixture(known=False))
        stats = excel.validate_excel(self.path)
        self.assertEqual(stats.collection_mode, "full_topsales")
        wb = load_workbook(self.path)
        try:
            self.assertEqual(tuple(wb["商品汇总"].cell(6, c).value for c in range(1, 15)), excel.SUMMARY_HEADERS)
            self.assertEqual(wb["抓取核验"].max_row, 23)
            self.assertEqual(wb["商品汇总"]["J7"].value, "未展示")
            self.assertEqual(wb["商品汇总"]["B7"].value, 1)
        finally:
            wb.close()

    def test_known_export_rejects_any_old_list_data(self):
        for field, value in {"global_rank": 1, "source_page": 0, "source_position": 1,
                             "price_php": 0, "monthly_sales_text": "0 Sold/Month",
                             "monthly_sales_count_lower_bound": 0,
                             "monthly_sales_observation": "not_displayed"}.items():
            with self.subTest(field=field):
                result = fixture()
                result["products"][0][field] = value
                with self.assertRaises(ValueError):
                    self.publish(result)
                self.assertFalse(self.path.exists())

    def test_known_export_rejects_invalid_scope(self):
        changes = {
            "reference_workbook": "relative.xlsx", "reference_sha256": "bad",
            "reference_captured_at": "2000/02/30 00:00:00", "reference_date": "2000-01-03",
            "reference_unique_count": 1, "reference_identity_sha256": "b" * 64,
            "user_confirmed_no_new_products": False,
        }
        for field, value in changes.items():
            with self.subTest(field=field):
                result = fixture()
                result["scope_metadata"][field] = value
                with self.assertRaises(ValueError):
                    self.publish(result)
                self.assertFalse(self.path.exists())

    def test_known_export_rejects_missing_or_wrong_required_audit(self):
        for field in fixture()["audit"]:
            for remove in (True, False):
                with self.subTest(field=field, remove=remove):
                    result = fixture()
                    if remove:
                        del result["audit"][field]
                    else:
                        result["audit"][field] += 1
                    with self.assertRaises(ValueError):
                        self.publish(result)

    def test_known_export_does_not_publish_partial_or_reordered_products(self):
        for change in (lambda result: result["products"].pop(),
                       lambda result: result["products"].reverse()):
            result = fixture()
            change(result)
            with self.assertRaises(ValueError):
                self.publish(result)

    def test_known_validator_rejects_old_fields_in_every_sheet(self):
        changes = [("商品汇总", cell, value) for cell, value in
                   (("B7", 1), ("C7", 1), ("G7", 0), ("H7", "0 Sold/Month"),
                    ("I7", 0), ("J7", "未展示"))]
        changes += [("SKU明细", "B7", 1), ("图片明细", "B7", 1)]
        for sheet, cell, value in changes:
            with self.subTest(sheet=sheet, cell=cell):
                self.publish()
                self.mutate(lambda wb: setattr(wb[sheet][cell], "value", value))
                with self.assertRaises(ValueError):
                    excel.validate_excel(self.path)

    def test_known_validator_rejects_scope_metadata_tampering(self):
        changes = {"参考工作簿": "/private/reference.xlsx", "参考日期": "2000-01-03",
                   "参考采集时间": "2001/01/01 00:00:00", "参考文件 SHA-256": "bad",
                   "参考身份序列 SHA-256": "b" * 64, "参考唯一商品数": 1,
                   "用户确认无新增商品": "否", "本次重抓列表": "是", "排序方式": "Top Sales",
                   "列表总页数": 1, "商品卡出现次数": 2, "跨页重复次数": 1,
                   "成功进入详情页": 1, "详情页失败数": 1, "未采集月销商品数": 0,
                   "完整性结论": "未通过"}
        for label, value in changes.items():
            with self.subTest(label=label):
                self.publish()
                self.mutate(lambda wb: setattr(wb["抓取核验"].cell(self.audit_row(wb, label), 2), "value", value))
                with self.assertRaises(ValueError):
                    excel.validate_excel(self.path)

    def test_known_validator_requires_headers_titles_and_scope(self):
        for sheet, coordinate, value in (("商品汇总", "A6", "全店排名"),
                                          ("SKU明细", "A6", "全店排名"),
                                          ("图片明细", "A6", "全店排名"),
                                          ("商品汇总", "A2", "UGREEN Shopee Top Sales 全店商品"),
                                          ("抓取核验", "B4", "UGREEN Top Sales 全部分页与全部详情页")):
            with self.subTest(sheet=sheet, cell=coordinate):
                self.publish()
                self.mutate(lambda wb: setattr(wb[sheet][coordinate], "value", value))
                with self.assertRaises(ValueError):
                    excel.validate_excel(self.path)

    def test_changing_mode_only_does_not_bypass_full_validation(self):
        self.publish(fixture(known=False))
        def change(wb):
            ws = wb["抓取核验"]
            ws.cell(8, 1).value = "采集模式"
            ws.cell(8, 2).value = "known_product_details"
        self.mutate(change)
        with self.assertRaises(ValueError):
            excel.validate_excel(self.path)

    def test_known_validator_keeps_identity_sku_gallery_and_formula_checks(self):
        for sheet, coordinate, value in (("商品汇总", "N7", "https://shopee.ph/product/64922227/999"),
                                          ("SKU明细", "D7", "999"), ("SKU明细", "F7", "not-numeric"),
                                          ("图片明细", "H7", "https://example.invalid/wrong.png"),
                                          ("图片明细", "G8", 2), ("商品汇总", "F7", "=1+1")):
            with self.subTest(sheet=sheet, cell=coordinate):
                self.publish()
                self.mutate(lambda wb: setattr(wb[sheet][coordinate], "value", value))
                with self.assertRaises(ValueError):
                    excel.validate_excel(self.path)

    def test_invalid_export_preserves_existing_file_and_cleans_temporary(self):
        self.publish()
        original = self.path.read_bytes()
        result = deepcopy(fixture())
        result["products"][0]["skus"] = []
        with self.assertRaises(ValueError):
            self.publish(result)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.directory.iterdir()), [self.path])


if __name__ == "__main__":
    unittest.main()
